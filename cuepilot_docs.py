"""
CuePilot reference documents.

Load a document (PDF, Word, or text), split it into passages, and when a question comes up
find the passage that best answers it and show it verbatim. Matching has three levels, best
to worst, and falls back automatically:

  * Contextual embedding search: each passage is embedded together with a short line saying
    where it sits in the document, so an isolated paragraph still matches questions about its
    surrounding topic. Needs an embeddings API (OpenAI or Gemini are cheap) and, for the
    context line, the AI you already set up. Done once on load.
  * Plain embedding search: same, without the context line, if the AI for context isn't set.
  * Keyword search: word overlap between the question and each passage. Free, offline, instant.

The app always shows YOUR document's own words. Search only decides which passage to show.
This module has no Qt dependency, so it can be unit-tested on its own.
"""
from __future__ import annotations

BUILD = "2026.10.21-7"  # must match APP_BUILD in cuepilot.py

import json
import math
import os
import re
import threading
import urllib.error
import urllib.parse
import urllib.request
from pathlib import Path

# Passage sizing, in characters (roughly 150-350 words per passage).
CHUNK_TARGET = 1100
CHUNK_MAX = 1800
CHUNK_OVERLAP = 150

# A result counts as a real answer when the top passage clearly stands out. Keyword and cosine
# scores live on different scales, so each has a floor; keyword also needs the question's words to
# actually appear. Below these, the document is judged not to contain the answer.
KEYWORD_MIN_SCORE = 0.02
SEMANTIC_MIN_SCORE = 0.30
KEYWORD_MIN_TERM_HITS = 1  # at least this many of the question's meaningful words must appear

STOPWORDS = {
    "the", "a", "an", "and", "or", "but", "if", "of", "to", "in", "on", "for", "with", "as",
    "at", "by", "from", "is", "are", "was", "were", "be", "been", "being", "it", "its", "this",
    "that", "these", "those", "do", "does", "did", "have", "has", "had", "will", "would",
    "should", "could", "can", "may", "might", "i", "you", "we", "they", "he", "she", "them",
    "us", "our", "your", "their", "my", "me", "what", "which", "who", "how", "when", "where",
    "why", "about", "into", "over", "than", "then", "so", "such", "not", "no", "yes", "please",
    "tell", "explain", "describe", "walk", "give",
}

EMBED_PROVIDERS = {
    # Uses the same Google sign-in as Claude on Google Cloud; billed to your Google Cloud project.
    "Google Cloud (Vertex AI)": {"kind": "vertex", "base_url": "", "model": "gemini-embedding-001",
                                 "env": "", "region": "us-central1"},
    # Anthropic doesn't make an embeddings model; its docs recommend Voyage AI.
    "Voyage AI (recommended by Anthropic)": {"base_url": "https://api.voyageai.com/v1",
                                             "model": "voyage-4-lite", "env": "VOYAGE_API_KEY"},
    "OpenAI": {"base_url": "https://api.openai.com/v1", "model": "text-embedding-3-small",
               "env": "OPENAI_API_KEY"},
    # text-embedding-004 was retired in January 2026; gemini-embedding-001 replaced it.
    "Google Gemini (free tier, no card)": {
        "base_url": "https://generativelanguage.googleapis.com/v1beta/openai",
        "model": "gemini-embedding-001", "env": "GEMINI_API_KEY"},
    # Free and local: install Ollama, then run 'ollama pull nomic-embed-text'. No account or key.
    "Ollama (local, free)": {"base_url": "http://localhost:11434/v1", "model": "nomic-embed-text",
                             "env": "", "needs_key": False},
    "Custom (OpenAI-compatible)": {"base_url": "", "model": "", "env": "", "needs_key": False},
}


# =============================================================================
# Reading files
# =============================================================================
class DocError(RuntimeError):
    pass


def read_document(path: str) -> str:
    p = Path(path)
    if not p.exists():
        raise DocError(f"The file wasn't found: {path}")
    ext = p.suffix.lower()
    if ext in (".txt", ".md", ".markdown", ".text"):
        return p.read_text(encoding="utf-8", errors="replace")
    if ext == ".pdf":
        return _read_pdf(p)
    if ext in (".docx",):
        return _read_docx(p)
    if ext == ".doc":
        raise DocError("Old .doc files aren't supported. Save it as .docx or PDF and try again.")
    raise DocError(f"Unsupported file type '{ext}'. Use PDF, Word (.docx), or a text file.")


def _read_pdf(p: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError:
        raise DocError("Reading PDFs needs the 'pypdf' package: pip install pypdf") from None
    try:
        reader = PdfReader(str(p))
    except Exception as exc:
        raise DocError(f"Couldn't open the PDF: {exc}") from None
    if getattr(reader, "is_encrypted", False):
        try:
            reader.decrypt("")
        except Exception:
            raise DocError("The PDF is password-protected. Remove the password and try again.") from None
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:
            pages.append("")
    text = "\n\n".join(pages).strip()
    if not text:
        raise DocError("No text could be read from this PDF. It may be a scan or images. "
                       "A text-based PDF or a Word file works best.")
    return text


def _read_docx(p: Path) -> str:
    try:
        import docx
    except ImportError:
        raise DocError("Reading Word files needs the 'python-docx' package: "
                       "pip install python-docx") from None
    try:
        doc = docx.Document(str(p))
    except Exception as exc:
        raise DocError(f"Couldn't open the Word file: {exc}") from None
    parts = [para.text for para in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    if not text:
        raise DocError("No text could be read from this Word file.")
    return text


# =============================================================================
# Chunking
# =============================================================================
_HEADING_RE = re.compile(r"^(#{1,6}\s+.+|[A-Z0-9][^\n]{0,70})$")


def _looks_like_heading(line: str) -> bool:
    s = line.strip()
    if not s or len(s) > 80:
        return False
    if s.startswith("#"):
        return True
    words = s.split()
    if len(words) <= 12 and not s.endswith((".", ",", ";", ":")):
        letters = [c for c in s if c.isalpha()]
        if letters and sum(c.isupper() for c in letters) / len(letters) > 0.6:
            return True
        if s.istitle() and len(words) <= 10:
            return True
    return False


def split_into_passages(text: str):
    """Return a list of {text, heading, index}. Splits on blank lines, keeps the nearest
    heading above each passage, and keeps passages within a size band."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = text.split("\n")
    blocks = []
    current = []
    heading = ""
    pending_heading = ""
    for line in lines:
        if _looks_like_heading(line) and not current:
            pending_heading = line.strip().lstrip("#").strip()
            heading = pending_heading
            continue
        if line.strip() == "":
            if current:
                blocks.append((heading, "\n".join(current).strip()))
                current = []
            continue
        if _looks_like_heading(line) and current:
            blocks.append((heading, "\n".join(current).strip()))
            current = []
            heading = line.strip().lstrip("#").strip()
            continue
        current.append(line)
    if current:
        blocks.append((heading, "\n".join(current).strip()))

    # Merge tiny blocks forward and split oversized ones, keeping an overlap.
    passages = []
    buf_head = ""
    buf = ""
    for head, body in blocks:
        if not body:
            continue
        if len(body) > CHUNK_MAX:
            if buf:
                passages.append((buf_head, buf))
                buf_head, buf = "", ""
            for piece in _split_long(body):
                passages.append((head, piece))
            continue
        if not buf:
            buf_head, buf = head, body
        elif head == buf_head and len(buf) + len(body) + 2 <= CHUNK_TARGET:
            # Only merge adjacent blocks that share a heading (or both have none).
            buf = buf + "\n\n" + body
        else:
            passages.append((buf_head, buf))
            buf_head, buf = head, body
    if buf:
        passages.append((buf_head, buf))

    return [{"index": i, "heading": h, "text": t} for i, (h, t) in enumerate(passages) if t.strip()]


def _split_long(body: str):
    """Split an oversized block on sentence boundaries into CHUNK_TARGET-sized pieces."""
    sentences = re.split(r"(?<=[.!?])\s+", body)
    pieces = []
    cur = ""
    for sent in sentences:
        if len(cur) + len(sent) + 1 <= CHUNK_TARGET or not cur:
            cur = (cur + " " + sent).strip()
        else:
            pieces.append(cur)
            tail = cur[-CHUNK_OVERLAP:]
            cur = (tail + " " + sent).strip()
    if cur:
        pieces.append(cur)
    return pieces


# =============================================================================
# Keyword search (always available)
# =============================================================================
def _tokens(text: str):
    return [w for w in re.findall(r"[a-z0-9']+", text.lower()) if w not in STOPWORDS and len(w) > 1]


def keyword_scores(question: str, passages):
    """Score each passage by weighted word overlap (a light TF-IDF). Returns a list of floats."""
    q_terms = set(_tokens(question))
    if not q_terms:
        return [0.0] * len(passages)
    # document frequency for inverse weighting
    df = {}
    passage_tokens = []
    for p in passages:
        toks = _tokens(p["text"] + " " + p.get("heading", ""))
        passage_tokens.append(toks)
        for t in set(toks):
            df[t] = df.get(t, 0) + 1
    n = max(1, len(passages))
    scores = []
    for toks in passage_tokens:
        if not toks:
            scores.append(0.0)
            continue
        counts = {}
        for t in toks:
            counts[t] = counts.get(t, 0) + 1
        score = 0.0
        for term in q_terms:
            if term in counts:
                idf = math.log((n + 1) / (df.get(term, 0) + 0.5))
                score += (counts[term] / len(toks)) * idf
        scores.append(score)
    return scores


# =============================================================================
# Embeddings (optional, better matching)
# =============================================================================
RETRY_CODES = (429, 500, 502, 503, 504)


def _post_json(url, body, headers, timeout=60, retries=4):
    """POST JSON. On a rate limit (429) or a temporary server error, wait and retry, as Voyage's
    docs recommend: honour Retry-After if given, otherwise back off 1, 2, 4, 8 seconds."""
    import time as _time
    data = json.dumps(body).encode("utf-8")
    attempt = 0
    while True:
        try:
            return _post_json_once(url, data, headers, timeout)
        except _Retryable as r:
            if attempt >= retries:
                raise DocError(r.message) from None
            wait = r.retry_after if r.retry_after is not None else min(15.0, 2.0 ** attempt)
            _time.sleep(max(0.5, min(wait, 30.0)))
            attempt += 1


class _Retryable(Exception):
    def __init__(self, message, retry_after=None):
        super().__init__(message)
        self.message = message
        self.retry_after = retry_after


def _post_json_once(url, data, headers, timeout):
    req = urllib.request.Request(url, data=data, headers=headers, method="POST")
    try:
        with urllib.request.urlopen(req, timeout=timeout) as resp:
            return json.loads(resp.read().decode("utf-8"))
    except urllib.error.HTTPError as exc:
        detail = ""
        raw = ""
        try:
            raw = exc.read().decode("utf-8", "replace")
            payload = json.loads(raw)
            if isinstance(payload, dict):
                err = payload.get("error") or payload.get("detail") or payload.get("message") or ""
            else:
                err = payload
            detail = (err.get("message", "") if isinstance(err, dict) else str(err))[:300]
        except Exception:
            detail = re.sub(r"\s+", " ", re.sub(r"<[^>]+>", " ", raw)).strip()[:200]
        hint = {401: "check the embeddings API key",
                403: "the key was refused; check it, and that the server URL matches where the key "
                     "came from (Voyage keys starting 'al-' use https://ai.mongodb.com/v1)",
                404: "check the server URL and model name",
                429: "rate limit reached"}.get(exc.code, "")
        msg = f"Embeddings failed (HTTP {exc.code})"
        if detail:
            msg += f": {detail}"
        if hint:
            msg += f" ({hint})"
        if exc.code in RETRY_CODES:
            retry_after = None
            try:
                retry_after = float(exc.headers.get("Retry-After"))
            except Exception:
                pass
            raise _Retryable(msg, retry_after) from None
        raise DocError(msg) from None
    except urllib.error.URLError as exc:
        host = urllib.parse.urlsplit(url).netloc or url
        raise DocError(f"Couldn't reach the embeddings server at {host}: {exc.reason}.") from None
    except Exception as exc:
        raise DocError(f"Embeddings failed: {exc}") from None


EMBED_BATCH = 32  # small enough for free-tier per-minute token limits


def _is_voyage(base: str) -> bool:
    b = base.lower()
    return "voyageai.com" in b or "ai.mongodb.com" in b


def embed_texts(cfg: dict, texts, timeout=120, input_type=None, retries=4):
    """Return a list of vectors for the given texts via an OpenAI-style embeddings endpoint.
    input_type ('document' or 'query') is sent to Voyage, which matches better with it."""
    if cfg.get("kind") == "vertex":
        try:
            import cuepilot_google as gcp
        except Exception:
            raise DocError("Google Cloud support isn't available in this build.") from None
        try:
            return gcp.vertex_embed(cfg, texts, input_type=input_type, timeout=min(timeout, 60),
                                    retries=retries)
        except Exception as exc:
            raise DocError(str(exc)) from None
    base = cfg["base_url"].strip().rstrip("/")
    if not base:
        raise DocError("No embeddings server URL is set.")
    headers = {"Content-Type": "application/json", "User-Agent": "CuePilot"}
    if cfg.get("api_key"):
        headers["Authorization"] = "Bearer " + cfg["api_key"]
    out = []
    for i in range(0, len(texts), EMBED_BATCH):
        batch = texts[i:i + EMBED_BATCH]
        body = {"model": cfg["model"], "input": batch}
        if input_type and _is_voyage(base):
            body["input_type"] = input_type
        data = _post_json(base + "/embeddings", body, headers, timeout, retries)
        try:
            rows = sorted(data["data"], key=lambda d: d.get("index", 0))
            out.extend(r["embedding"] for r in rows)
        except (KeyError, TypeError):
            raise DocError("The embeddings server sent an unexpected reply. Check the model name.") from None
    return out


def _norm(vec):
    s = math.sqrt(sum(x * x for x in vec))
    return s or 1.0


def cosine(a, b):
    if not a or not b or len(a) != len(b):
        return 0.0
    dot = sum(x * y for x, y in zip(a, b))
    return dot / (_norm(a) * _norm(b))


def cosine_scores(query_vec, matrix):
    try:
        import numpy as np
        if not matrix:
            return []
        m = np.asarray(matrix, dtype=np.float32)
        q = np.asarray(query_vec, dtype=np.float32)
        mn = np.linalg.norm(m, axis=1)
        mn[mn == 0] = 1.0
        qn = np.linalg.norm(q) or 1.0
        return (m @ q / (mn * qn)).tolist()
    except Exception:
        return [cosine(query_vec, row) for row in matrix]


# =============================================================================
# The reference library
# =============================================================================
CONTEXT_SYSTEM = (
    "You label passages from a document so they can be searched. Given the whole document and one "
    "passage from it, write a single short sentence (under 25 words) saying what the passage is "
    "about and where it fits in the document. Output only that sentence, no preamble."
)


class ReferenceDoc:
    """One loaded document: its passages, optional embeddings, and search."""

    def __init__(self, name: str):
        self.name = name
        self.passages = []        # list of {index, heading, text, context}
        self.vectors = []         # parallel list of embedding vectors (may be empty)
        self.embedded = False
        self.contextual = False

    def load_text(self, text: str):
        self.passages = split_into_passages(text)
        self.vectors = []
        self.embedded = False
        self.contextual = False
        return len(self.passages)

    def build_embeddings(self, embed_cfg, context_ai=None, on_progress=None, should_stop=None):
        """Embed every passage. If context_ai is given, prepend a one-line context header first
        (contextual retrieval). Runs synchronously; call it on a worker thread."""
        if not self.passages:
            return
        full = "\n\n".join(p["text"] for p in self.passages)[:16000]
        contexts = self._build_contexts(full, context_ai, on_progress, should_stop)
        if contexts is None:
            return  # stopped
        to_embed = []
        for p, ctx in zip(self.passages, contexts):
            p["context"] = ctx
            header = (f"{p['heading']}. " if p["heading"] else "") + (ctx + " " if ctx else "")
            to_embed.append((header + p["text"]).strip())
        if should_stop and should_stop():
            return
        if on_progress:
            on_progress(len(self.passages), len(self.passages), "embedding")
        self.vectors = embed_texts(embed_cfg, to_embed, input_type="document")
        self.embedded = len(self.vectors) == len(self.passages)
        self.contextual = self.embedded and any(contexts)

    CONTEXT_WORKERS = 4
    CONTEXT_GIVE_UP_AFTER = 3  # consecutive failures before skipping context notes

    def _build_contexts(self, full, context_ai, on_progress, should_stop):
        """One-line context notes for every passage, written by the AI several at a time. If the
        AI keeps failing, stop asking and carry on without notes rather than stalling."""
        n = len(self.passages)
        if context_ai is None:
            return [""] * n
        from concurrent.futures import ThreadPoolExecutor, as_completed
        results = [""] * n
        failures = 0
        done = 0
        gave_up = threading.Event()

        def one(i):
            if gave_up.is_set() or (should_stop and should_stop()):
                return i, None
            return i, _passage_context(context_ai, full, self.passages[i]["text"])

        with ThreadPoolExecutor(max_workers=self.CONTEXT_WORKERS) as pool:
            futures = [pool.submit(one, i) for i in range(n)]
            for fut in as_completed(futures):
                i, ctx = fut.result()
                done += 1
                if ctx:
                    results[i] = ctx
                    failures = 0
                elif ctx == "":
                    failures += 1
                    if failures >= self.CONTEXT_GIVE_UP_AFTER:
                        gave_up.set()
                if on_progress:
                    on_progress(done, n, "preparing")
        if should_stop and should_stop():
            return None
        return results

    def search_keyword(self, question: str, top_k=1):
        """Instant, offline keyword search. No network, returns at once."""
        if not self.passages:
            return []
        kw = keyword_scores(question, self.passages)
        return self._rank(question, kw, "keyword", top_k)

    def search_semantic(self, question: str, embed_cfg, top_k=1):
        """Semantic (or contextual) search. One embedding call for the question. Raises DocError
        on network trouble, so the caller can keep the keyword result it already has."""
        if not self.passages or not self.embedded or embed_cfg is None:
            return self.search_keyword(question, top_k)
        kw = keyword_scores(question, self.passages)
        # Fast path for a live question: short timeout and no retries, so a slow moment falls back
        # to the keyword result that's already on screen instead of making you wait.
        qvec = embed_texts(embed_cfg, [question], timeout=8, input_type="query", retries=0)[0]
        sem = cosine_scores(qvec, self.vectors)
        if not sem or len(sem) != len(self.passages):
            return self._rank(question, kw, "keyword", top_k)
        kmax = max(kw) or 1.0
        scores = [0.8 * s + 0.2 * (k / kmax) for s, k in zip(sem, kw)]
        method = "contextual" if self.contextual else "semantic"
        return self._rank(question, scores, method, top_k)

    def search(self, question: str, embed_cfg=None, top_k=1):
        """Best available single-shot search: semantic if set up, else keyword."""
        if self.embedded and embed_cfg is not None:
            try:
                return self.search_semantic(question, embed_cfg, top_k)
            except DocError:
                pass
        return self.search_keyword(question, top_k)

    def _rank(self, question, scores, method, top_k):
        order = sorted(range(len(self.passages)), key=lambda i: scores[i], reverse=True)
        q_terms = set(_tokens(question))
        results = []
        for i in order[:top_k]:
            p = self.passages[i]
            results.append({"heading": p["heading"], "text": p["text"],
                            "score": scores[i], "method": method, "index": p["index"],
                            "good": self._is_good(scores[i], method, q_terms, p)})
        return results

    @staticmethod
    def _is_good(score, method, q_terms, passage):
        if method in ("semantic", "contextual"):
            return score >= SEMANTIC_MIN_SCORE
        # keyword: need a minimum score AND the question's words actually present
        hits = len(q_terms & set(_tokens(passage["text"] + " " + passage.get("heading", ""))))
        return score >= KEYWORD_MIN_SCORE and hits >= KEYWORD_MIN_TERM_HITS


def _passage_context(context_ai, full_doc: str, passage: str) -> str:
    from_app = context_ai["complete"]  # a callable (system, user, max_tokens) -> str; it should
    # use a short timeout of its own so one slow reply can't hold up the document
    user = (f"<document>\n{full_doc}\n</document>\n\n<passage>\n{passage[:1500]}\n</passage>\n\n"
            "Give the one-sentence context for the passage.")
    try:
        text = from_app(CONTEXT_SYSTEM, user, 80)
    except Exception:
        return ""
    return (text or "").strip().replace("\n", " ")[:200]


def build_embed_config(provider_name: str, url: str, key: str, model: str, project: str = "",
                       region: str = ""):
    meta = EMBED_PROVIDERS.get(provider_name, EMBED_PROVIDERS["Custom (OpenAI-compatible)"])
    if not model.strip():
        return None, "Enter an embeddings model name."
    if meta.get("kind") == "vertex":
        if not project.strip():
            return None, "Enter your Google Cloud project ID for document search."
        return {"kind": "vertex", "project": project.strip(), "model": model.strip(),
                "region": region.strip() or meta.get("region", "us-central1")}, None
    if not url.strip():
        return None, "Enter the embeddings server URL."
    if not key.strip() and meta.get("env"):
        key = os.environ.get(meta["env"], "")
    if not key.strip() and meta.get("needs_key", True):
        return None, f"Add an API key for {provider_name} embeddings."
    url = re.sub(r"/vi(?=/|$)", "/v1", url.strip().rstrip("/"), flags=re.IGNORECASE)
    if url.lower().endswith("/embeddings"):
        url = url[: -len("/embeddings")].rstrip("/")
    return {"base_url": voyage_host_for_key(url, key.strip()), "api_key": key.strip(),
            "model": model.strip()}, None


VOYAGE_DIRECT = "https://api.voyageai.com/v1"
VOYAGE_ATLAS = "https://ai.mongodb.com/v1"


def voyage_host_for_key(url: str, key: str) -> str:
    """Voyage is part of MongoDB: keys from MongoDB Atlas start with 'al-' and only work at
    ai.mongodb.com; keys from the Voyage dashboard start with 'pa-' and only work at
    api.voyageai.com. Each host refuses the other's keys with HTTP 403, so pick the right one."""
    host = url.rstrip("/").lower()
    if key.startswith("al-") and "api.voyageai.com" in host:
        return VOYAGE_ATLAS
    if key.startswith("pa-") and "ai.mongodb.com" in host:
        return VOYAGE_DIRECT
    return url
